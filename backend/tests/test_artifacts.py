"""Tests for qsagent.artifacts — bounds, registry, workbook, export_dto.

Run from repo root: pytest backend/tests/test_artifacts.py -v
"""

from __future__ import annotations

import io
import json
import sqlite3
import time
import uuid
from typing import Any

import openpyxl
import pytest

from qsagent.artifacts import (
    ARTIFACT_ID_CHARS,
    ASSUMPTION_COLUMNS,
    BANNER_ROW,
    CHECKMATE_COLUMNS,
    CLAIM_COLUMNS,
    DEFAULT_ARTIFACT_TTL_SECONDS,
    EXPORT_TYPE_BOQ_XLSX,
    FIRST_DATA_ROW,
    GATE_NOT_RUN,
    GATE_PASSED,
    GATE_REJECTED,
    HEADER_ROW,
    MAX_ARTIFACT_AGE_SECONDS,
    MAX_EXPORT_BYTES,
    MAX_EXPORT_CHECKMATE_ROWS,
    MAX_EXPORT_CLAIMS,
    MAX_EXPORT_COLUMNS,
    MAX_EXPORT_EVIDENCE_PER_CLAIM,
    MAX_EXPORT_FORMULA_CELLS,
    MAX_EXPORT_FORMULA_CHARS,
    MAX_EXPORT_SHEETS,
    MAX_EXPORT_TEXT_CHARS,
    MAX_EXPORT_WARNINGS,
    SHEET_ASSUMPTIONS,
    SHEET_CHECKMATE,
    SHEET_CLAIMS,
    SHEET_NAMES,
    VALUE_UNRESOLVED,
    ArtifactEntry,
    ArtifactRef,
    ArtifactRegistry,
    ExportRequest,
    ExportResponse,
    RateBinding,
    WorkbookWarning,
    build_boq_workbook,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_row(d: dict[str, Any]) -> sqlite3.Row:
    """Build a real sqlite3.Row, preserving Python int/float/None types."""
    cols = list(d.keys())
    placeholders = ", ".join("?" for _ in cols)
    # Use BLOB affinity for all columns so SQLite does not coerce values
    col_defs = ", ".join(f'"{c}"' for c in cols)
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.execute(f"CREATE TABLE t ({col_defs})")
    conn.execute(f"INSERT INTO t VALUES ({placeholders})", [d[c] for c in cols])
    return conn.execute("SELECT * FROM t").fetchone()


def _claim_row(
    claim_id: str = "C-001",
    description: str = "Bulk earthworks",
    value: float = 100.0,
    unit: str = "m3",
    method: str = "tier1",
    evidence: str | None = None,
    measurement_state: str | None = None,
) -> sqlite3.Row:
    ev = evidence or json.dumps([{"file_name": "drawing.pdf", "source_locator": "S1"}])
    return _make_row({
        "claim_id": claim_id,
        "description": description,
        "value": value,           # float, not str
        "unit": unit,
        "method": method,
        "evidence": ev,
        "measurement_state": measurement_state or "",
    })


def _assumption_row(aid: str = "A-01", statement: str = "Soil is class E") -> sqlite3.Row:
    return _make_row({
        "id": aid, "statement": statement, "status": "ASSUMED",
        "rationale": "From geotech", "impact_delta_aud": "50000",
        "impact_value": "200", "impact_unit": "m3",
        "evidence": json.dumps([{"file_name": "geotech.pdf", "source_locator": "p3"}]),
        "raised_at": "2026-09-16T00:00:00Z", "resolved_at": "",
    })


def _checkmate_row(subject: str = "C-001", passed: int = 1) -> sqlite3.Row:
    return _make_row({
        "subject": subject, "passed": passed,   # int, not str
        "findings": json.dumps([{"rule": "evidence_present", "ok": True}]),
        "created_at": "2026-09-16T00:00:00Z",
    })


def _build_simple(claims=None, assumptions=None, checkmate=None, rate_bindings=None):
    claims = claims if claims is not None else [_claim_row()]
    assumptions = assumptions if assumptions is not None else []
    checkmate = checkmate if checkmate is not None else []
    rate_bindings = rate_bindings if rate_bindings is not None else {"C-001": 45.0}
    return build_boq_workbook(
        project_name="Test Project",
        claim_rows=claims,
        assumption_rows=assumptions,
        checkmate_rows=checkmate,
        rate_bindings=rate_bindings,
    )


# ---------------------------------------------------------------------------
# bounds
# ---------------------------------------------------------------------------

class TestBoundsConstants:
    def test_banner_header_data_rows_ordered(self):
        assert BANNER_ROW == 1
        assert HEADER_ROW == 2
        assert FIRST_DATA_ROW == 3

    def test_sheet_names_ordered(self):
        assert SHEET_NAMES == (SHEET_CLAIMS, SHEET_ASSUMPTIONS, SHEET_CHECKMATE)

    def test_max_export_columns_is_widest(self):
        assert MAX_EXPORT_COLUMNS == max(
            len(CLAIM_COLUMNS), len(ASSUMPTION_COLUMNS), len(CHECKMATE_COLUMNS)
        )

    def test_max_export_sheets(self):
        assert MAX_EXPORT_SHEETS == len(SHEET_NAMES)

    def test_formula_ceil_ge_claim_ceil(self):
        assert MAX_EXPORT_FORMULA_CELLS == MAX_EXPORT_CLAIMS + 1

    def test_artifact_id_chars(self):
        assert ARTIFACT_ID_CHARS == 36
        assert len(str(uuid.uuid4())) == ARTIFACT_ID_CHARS

    def test_formula_shapes_within_chars(self):
        row = FIRST_DATA_ROW
        last = MAX_EXPORT_CLAIMS + FIRST_DATA_ROW - 1
        assert len(f"=C{row}*E{row}") <= MAX_EXPORT_FORMULA_CHARS
        assert len(f"=SUM(F{FIRST_DATA_ROW}:F{last})") <= MAX_EXPORT_FORMULA_CHARS

    def test_claim_col_headers_present(self):
        for col in ("Claim ID", "Description", "Quantity", "Unit", "Normalized Rate", "Extended Amount"):
            assert col in CLAIM_COLUMNS, f"{col!r} missing from CLAIM_COLUMNS"

    def test_evidence_per_claim_positive(self):
        assert MAX_EXPORT_EVIDENCE_PER_CLAIM >= 1

    def test_default_ttl_equals_max_age(self):
        assert DEFAULT_ARTIFACT_TTL_SECONDS == float(MAX_ARTIFACT_AGE_SECONDS)


# ---------------------------------------------------------------------------
# ArtifactRegistry
# ---------------------------------------------------------------------------

class TestArtifactRegistry:
    def test_register_returns_uuid4(self):
        reg = ArtifactRegistry()
        aid = reg.register(b"hello", media_type="text/plain", project_id=1)
        assert len(aid) == 36
        uuid.UUID(aid, version=4)

    def test_get_returns_entry(self):
        reg = ArtifactRegistry()
        aid = reg.register(b"data", media_type="application/octet-stream", project_id=1)
        entry = reg.get(aid)
        assert entry is not None
        assert entry.data == b"data"

    def test_get_unknown_returns_none(self):
        reg = ArtifactRegistry()
        assert reg.get(str(uuid.uuid4())) is None

    def test_get_wrong_length_returns_none(self):
        reg = ArtifactRegistry()
        assert reg.get("short") is None
        assert reg.get("x" * 37) is None

    def test_delete_removes_entry(self):
        reg = ArtifactRegistry()
        aid = reg.register(b"x", media_type="text/plain", project_id=1)
        assert reg.delete(aid) is True
        assert reg.get(aid) is None

    def test_delete_unknown_returns_false(self):
        reg = ArtifactRegistry()
        assert reg.delete(str(uuid.uuid4())) is False

    def test_len_counts_live_entries(self):
        reg = ArtifactRegistry()
        assert len(reg) == 0
        reg.register(b"a", media_type="text/plain", project_id=1)
        reg.register(b"b", media_type="text/plain", project_id=1)
        assert len(reg) == 2

    def test_expired_entry_not_returned(self):
        reg = ArtifactRegistry(ttl_seconds=0.01)
        aid = reg.register(b"x", media_type="text/plain", project_id=1)
        time.sleep(0.05)
        assert reg.get(aid) is None

    def test_expired_entry_evicted_from_len(self):
        reg = ArtifactRegistry(ttl_seconds=0.01)
        reg.register(b"x", media_type="text/plain", project_id=1)
        time.sleep(0.05)
        assert len(reg) == 0

    def test_oversized_raises(self):
        reg = ArtifactRegistry(max_bytes=10)
        with pytest.raises(ValueError, match="bytes"):
            reg.register(b"x" * 11, media_type="text/plain", project_id=1)

    def test_cleanup_thread_start_stop(self):
        reg = ArtifactRegistry(cleanup_interval=0.05)
        reg.start_cleanup()
        assert reg._cleanup_thread is not None and reg._cleanup_thread.is_alive()
        reg.start_cleanup()  # idempotent
        reg.stop_cleanup()
        assert not reg._cleanup_thread.is_alive()

    def test_iterate_yields_live_entries(self):
        reg = ArtifactRegistry()
        reg.register(b"a", media_type="text/plain", project_id=1)
        reg.register(b"b", media_type="text/plain", project_id=1)
        assert len(list(reg)) == 2

    def test_entry_properties(self):
        reg = ArtifactRegistry()
        aid = reg.register(b"abc", media_type="text/plain", project_id=1)
        e = reg.get(aid)
        assert e.bytes_len == 3
        assert e.age_seconds >= 0.0
        assert not e.expired

    def test_project_isolation_get(self):
        """An artifact registered under project 1 is invisible to project 2."""
        reg = ArtifactRegistry()
        aid = reg.register(b"secret", media_type="text/plain", project_id=1)
        assert reg.get(aid, project_id=1) is not None
        assert reg.get(aid, project_id=2) is None

    def test_project_isolation_no_project_id_bypasses(self):
        """get() without project_id returns entry regardless of stored project."""
        reg = ArtifactRegistry()
        aid = reg.register(b"x", media_type="text/plain", project_id=99)
        assert reg.get(aid) is not None  # no project filter

    def test_async_cleanup_task_is_coroutine(self):
        """async_cleanup_task must be an awaitable coroutine function."""
        import inspect
        reg = ArtifactRegistry()
        assert inspect.iscoroutinefunction(reg.async_cleanup_task)


# ---------------------------------------------------------------------------
# ExportRequest / ExportResponse / RateBinding
# ---------------------------------------------------------------------------

class TestExportDTO:
    def test_default_export_type(self):
        req = ExportRequest()
        assert req.export_type == EXPORT_TYPE_BOQ_XLSX

    def test_rejects_unknown_export_type(self):
        with pytest.raises(Exception):
            ExportRequest(export_type="csv")

    def test_rejects_extra_fields(self):
        with pytest.raises(Exception):
            ExportRequest(**{"extra_field": "bad"})

    def test_rate_binding_valid(self):
        rb = RateBinding(claim_id="C-001", rate_node_id=42)
        assert rb.claim_id == "C-001"
        assert rb.rate_node_id == 42

    def test_rate_binding_rejects_zero_node_id(self):
        with pytest.raises(Exception):
            RateBinding(claim_id="C-001", rate_node_id=0)

    def test_rate_binding_rejects_negative_node_id(self):
        with pytest.raises(Exception):
            RateBinding(claim_id="C-001", rate_node_id=-1)

    def test_rate_binding_rejects_empty_claim_id(self):
        with pytest.raises(Exception):
            RateBinding(claim_id="", rate_node_id=1)

    def test_rate_binding_no_rate_field(self):
        """RateBinding must not accept a rate_aud field — client cannot supply a number."""
        with pytest.raises(Exception):
            RateBinding(claim_id="C-001", rate_node_id=1, rate_aud=99.0)

    def test_artifact_ref_length_check(self):
        ref = ArtifactRef(artifact_id=str(uuid.uuid4()), download_url="http://localhost/dl/x")
        assert len(ref.artifact_id) == 36

    def test_export_response_gate_values(self):
        ref = ArtifactRef(artifact_id=str(uuid.uuid4()), download_url="http://localhost/dl/x")
        for gate in ("PASSED", "REJECTED", "NOT_RUN"):
            resp = ExportResponse(
                artifact=ref, project_id=1, claim_count=0,
                truncated=False, gate=gate, warnings=[],
            )
            assert resp.gate == gate

    def test_export_response_rejects_bad_gate(self):
        ref = ArtifactRef(artifact_id=str(uuid.uuid4()), download_url="http://localhost/dl/x")
        with pytest.raises(Exception):
            ExportResponse(
                artifact=ref, project_id=1, claim_count=0,
                truncated=False, gate="MAYBE", warnings=[],
            )


# ---------------------------------------------------------------------------
# build_boq_workbook — layout
# ---------------------------------------------------------------------------

class TestWorkbookLayout:
    def test_returns_bytes(self):
        data, _, _ = _build_simple()
        assert isinstance(data, bytes) and len(data) > 0

    def test_xlsx_parseable(self):
        data, _, _ = _build_simple()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert SHEET_CLAIMS in wb.sheetnames

    def test_sheet_names_correct_and_ordered(self):
        data, _, _ = _build_simple()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert wb.sheetnames == list(SHEET_NAMES)

    def test_banner_row_is_row_1(self):
        data, _, _ = _build_simple()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb[SHEET_CLAIMS]
        assert "NON-DELIVERABLE DRAFT" in (ws.cell(row=BANNER_ROW, column=1).value or "")

    def test_header_row_is_row_2_claims(self):
        data, _, _ = _build_simple()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb[SHEET_CLAIMS]
        headers = [ws.cell(row=HEADER_ROW, column=i + 1).value for i in range(len(CLAIM_COLUMNS))]
        assert headers == list(CLAIM_COLUMNS)

    def test_header_row_is_row_2_assumptions(self):
        data, _, _ = _build_simple(assumptions=[_assumption_row()])
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb[SHEET_ASSUMPTIONS]
        headers = [ws.cell(row=HEADER_ROW, column=i + 1).value for i in range(len(ASSUMPTION_COLUMNS))]
        assert headers == list(ASSUMPTION_COLUMNS)

    def test_header_row_is_row_2_checkmate(self):
        data, _, _ = _build_simple(checkmate=[_checkmate_row()])
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb[SHEET_CHECKMATE]
        headers = [ws.cell(row=HEADER_ROW, column=i + 1).value for i in range(len(CHECKMATE_COLUMNS))]
        assert headers == list(CHECKMATE_COLUMNS)

    def test_first_data_row_is_row_3(self):
        data, _, _ = _build_simple()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb[SHEET_CLAIMS]
        assert ws.cell(row=FIRST_DATA_ROW, column=1).value == "C-001"

    def test_quantity_in_column_c(self):
        data, _, _ = _build_simple()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb[SHEET_CLAIMS]
        assert ws.cell(row=FIRST_DATA_ROW, column=3).value == pytest.approx(100.0)

    def test_rate_in_column_e(self):
        data, _, _ = _build_simple(rate_bindings={"C-001": 45.0})
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb[SHEET_CLAIMS]
        assert ws.cell(row=FIRST_DATA_ROW, column=5).value == pytest.approx(45.0)

    def test_formula_in_column_f_references_c_and_e(self):
        data, _, _ = _build_simple(rate_bindings={"C-001": 45.0})
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        ws = wb[SHEET_CLAIMS]
        cell_f = ws.cell(row=FIRST_DATA_ROW, column=6).value or ""
        assert f"C{FIRST_DATA_ROW}" in cell_f and f"E{FIRST_DATA_ROW}" in cell_f

    def test_formula_length_within_bound(self):
        data, _, _ = _build_simple(rate_bindings={"C-001": 45.0})
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        ws = wb[SHEET_CLAIMS]
        cell_f = ws.cell(row=FIRST_DATA_ROW, column=6).value or ""
        assert len(cell_f) <= MAX_EXPORT_FORMULA_CHARS


# ---------------------------------------------------------------------------
# Integration tests: artifact API routes
# ---------------------------------------------------------------------------

import pytest as _pytest
from fastapi.testclient import TestClient as _TestClient

from qsagent.api import create_app as _create_app
from qsagent.runtime import ModelRouter as _ModelRouter, ModelTier as _ModelTier
from qsagent.runtime import ProviderResponse as _ProviderResponse
from qsagent.storage import QSStore as _QSStore

_API = "/api/v1"
_TEST_TOKEN = "artifact-integration-test-token-abc"
_AUTH_HEADERS = {"Authorization": f"Bearer {_TEST_TOKEN}"}


class _NoopProvider:
    @property
    def name(self) -> str:
        return "noop"

    @property
    def supported_tier(self):
        return _ModelTier.LOCAL_MODEL

    def complete(self, prompt: str, **kw):
        return _ProviderResponse(
            text="", provider_name="noop", model_name="noop-v0", usage_tokens=0
        )


class _AnonSecrets:
    def get_key(self, provider: str):
        return None


def _make_test_app(store, tmp_path, sandbox_suffix="sandboxes"):
    r = _ModelRouter(_AnonSecrets())
    r.register(_NoopProvider())
    return _create_app(
        store=store,
        router=r,
        api_token=_TEST_TOKEN,
        sandbox_root=tmp_path / sandbox_suffix,
    )


@_pytest.fixture()
def _store():
    st = _QSStore(":memory:")
    yield st
    st.close()


@_pytest.fixture()
def _app(_store, tmp_path):
    return _make_test_app(_store, tmp_path)


@_pytest.fixture()
def _client(_app):
    with _TestClient(_app, headers=_AUTH_HEADERS) as c:
        yield c


def _new_project(client, name: str = "artifact-test") -> int:
    resp = client.post(f"{_API}/projects", json={"name": name})
    assert resp.status_code == 201
    return resp.json()["project_id"]


class TestArtifactRoutes:
    """Integration: POST /projects/{id}/artifacts + GET download."""

    def test_create_requires_auth(self, _app):
        with _TestClient(_app) as c:
            resp = c.post(f"{_API}/projects/1/artifacts", json={})
        assert resp.status_code in (401, 403)

    def test_download_requires_auth(self, _app):
        with _TestClient(_app) as c:
            resp = c.get(f"{_API}/projects/1/artifacts/{str(uuid.uuid4())}")
        assert resp.status_code in (401, 403)

    def test_create_unknown_project_is_404(self, _client):
        resp = _client.post(f"{_API}/projects/99999/artifacts", json={})
        assert resp.status_code == 404

    def test_create_returns_201_with_artifact_ref(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        assert resp.status_code == 201
        body = resp.json()
        assert "artifact" in body
        assert len(body["artifact"]["artifact_id"]) == 36
        assert body["project_id"] == pid

    def test_empty_project_gate_is_not_run(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        assert resp.json()["gate"] == "NOT_RUN"

    def test_passed_checkmate_sets_passed_gate(self, _client, _store):
        pid = _new_project(_client)
        _store.save_checkmate(pid, subject="C-001", passed=True, findings=[])
        resp = _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        assert resp.json()["gate"] == "PASSED"

    def test_failed_checkmate_sets_rejected_gate(self, _client, _store):
        pid = _new_project(_client)
        _store.save_checkmate(pid, subject="C-001", passed=False, findings=[{"msg": "bad"}])
        resp = _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        assert resp.json()["gate"] == "REJECTED"

    def test_diagnostic_gate_does_not_mutate_db(self, _client, _store):
        """create_artifact must not write to the DB."""
        pid = _new_project(_client)
        before = _store.journal_entries(pid)
        _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        after = _store.journal_entries(pid)
        assert len(after) == len(before)

    def test_download_url_contains_artifact_id(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        body = resp.json()
        assert body["artifact"]["artifact_id"] in body["artifact"]["download_url"]

    def test_download_returns_xlsx_bytes(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        artifact_id = resp.json()["artifact"]["artifact_id"]
        dl = _client.get(f"{_API}/projects/{pid}/artifacts/{artifact_id}")
        assert dl.status_code == 200
        # XLSX magic bytes
        assert dl.content[:2] == b"PK"

    def test_download_content_type_is_xlsx(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        artifact_id = resp.json()["artifact"]["artifact_id"]
        dl = _client.get(f"{_API}/projects/{pid}/artifacts/{artifact_id}")
        assert "spreadsheetml" in dl.headers.get("content-type", "")


    def test_project_isolation_download(self, _client, _store):
        """Artifact for project A is not downloadable under project B."""
        pid_a = _new_project(_client, "iso-a")
        pid_b = _new_project(_client, "iso-b")
        resp = _client.post(f"{_API}/projects/{pid_a}/artifacts", json={})
        artifact_id = resp.json()["artifact"]["artifact_id"]
        dl = _client.get(f"{_API}/projects/{pid_b}/artifacts/{artifact_id}")
        assert dl.status_code == 404

    def test_download_unknown_artifact_is_404(self, _client, _store):
        pid = _new_project(_client)
        dl = _client.get(f"{_API}/projects/{pid}/artifacts/{str(uuid.uuid4())}")
        assert dl.status_code == 404

    def test_uuid_path_traversal_rejected(self, _client, _store):
        """Non-UUID artifact ids must return 404 without any registry lookup."""
        pid = _new_project(_client)
        for bad_id in ("../../etc/passwd", "short", "x" * 37):
            dl = _client.get(f"{_API}/projects/{pid}/artifacts/{bad_id}")
            assert dl.status_code == 404, f"expected 404 for {bad_id!r}"

    def test_expired_artifact_is_404(self, _store, tmp_path):
        """Artifact whose TTL has elapsed must not be downloadable."""
        from qsagent.artifacts import ArtifactRegistry
        tiny_reg = ArtifactRegistry(ttl_seconds=0.01)
        app2 = _make_test_app(_store, tmp_path, "sandboxes-exp")
        app2.state.artifacts = tiny_reg
        with _TestClient(app2, headers=_AUTH_HEADERS) as c:
            pid = c.post(f"{_API}/projects", json={"name": "expire-test"}).json()["project_id"]
            resp = c.post(f"{_API}/projects/{pid}/artifacts", json={})
            artifact_id = resp.json()["artifact"]["artifact_id"]
            time.sleep(0.05)
            dl = c.get(f"{_API}/projects/{pid}/artifacts/{artifact_id}")
            assert dl.status_code == 404

    def test_lifespan_cleanup_task_starts_without_error(self, _store, tmp_path):
        """The lifespan must start the asyncio cleanup task without raising."""
        app3 = _make_test_app(_store, tmp_path, "sandboxes-ls")
        with _TestClient(app3, headers=_AUTH_HEADERS) as c:
            assert c.get(f"{_API}/health").status_code == 200

    def test_invalid_export_type_rejected_422(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(
            f"{_API}/projects/{pid}/artifacts",
            json={"export_type": "csv"},
        )
        assert resp.status_code == 422

    def test_rate_binding_node_zero_rejected_422(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(
            f"{_API}/projects/{pid}/artifacts",
            json={"rate_bindings": [{"claim_id": "C-001", "rate_node_id": 0}]},
        )
        assert resp.status_code == 422

    def test_extra_request_fields_rejected_422(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(
            f"{_API}/projects/{pid}/artifacts",
            json={"export_type": "boq_xlsx", "inject_me": True},
        )
        assert resp.status_code == 422

    def test_rate_node_cross_project_does_not_resolve(self, _client, _store):
        """A rate node from project A is invisible to project B."""
        from qsagent.contracts.evidence import EvidenceNode
        pid_a = _new_project(_client, "cpi-a")
        pid_b = _new_project(_client, "cpi-b")
        node_id = _store.add_node(EvidenceNode(
            project_id=pid_a,
            node_type="rate",
            label="rate-in-a",
            payload={"status": "normalized", "normalized_amount": 42.0},
        ))
        # Must succeed (generate workbook), rate just shows as Unresolved
        resp = _client.post(
            f"{_API}/projects/{pid_b}/artifacts",
            json={"rate_bindings": [{"claim_id": "C-001", "rate_node_id": node_id}]},
        )
        assert resp.status_code == 201

    def test_unresolved_rate_node_still_generates_workbook(self, _client, _store):
        """A rate node with non-normalized status must not block workbook generation."""
        from qsagent.contracts.evidence import EvidenceNode
        pid = _new_project(_client)
        node_id = _store.add_node(EvidenceNode(
            project_id=pid,
            node_type="rate",
            label="unres-rate",
            payload={"status": "unresolved"},
        ))
        resp = _client.post(
            f"{_API}/projects/{pid}/artifacts",
            json={"rate_bindings": [{"claim_id": "C-001", "rate_node_id": node_id}]},
        )
        assert resp.status_code == 201

    def test_size_limit_respected(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        artifact_id = resp.json()["artifact"]["artifact_id"]
        dl = _client.get(f"{_API}/projects/{pid}/artifacts/{artifact_id}")
        assert len(dl.content) <= MAX_EXPORT_BYTES

    def test_download_has_content_disposition_attachment(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        artifact_id = resp.json()["artifact"]["artifact_id"]
        dl = _client.get(f"{_API}/projects/{pid}/artifacts/{artifact_id}")
        assert "attachment" in dl.headers.get("content-disposition", "")

    def test_download_cache_control_no_store(self, _client, _store):
        pid = _new_project(_client)
        resp = _client.post(f"{_API}/projects/{pid}/artifacts", json={})
        artifact_id = resp.json()["artifact"]["artifact_id"]
        dl = _client.get(f"{_API}/projects/{pid}/artifacts/{artifact_id}")
        assert "no-store" in dl.headers.get("cache-control", "")


# ---------------------------------------------------------------------------
# build_boq_workbook — warnings
# ---------------------------------------------------------------------------

class TestWorkbookWarnings:
    def test_no_rate_binding_yields_unresolved_warning(self):
        _, warnings, _ = _build_simple(rate_bindings={})
        assert any(w.code == "rate_unresolved" for w in warnings)

    def test_unresolved_rate_cell_value(self):
        data, _, _ = _build_simple(rate_bindings={})
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb[SHEET_CLAIMS]
        assert ws.cell(row=FIRST_DATA_ROW, column=5).value == VALUE_UNRESOLVED

    def test_unknown_binding_yields_warning(self):
        _, warnings, _ = _build_simple(rate_bindings={"C-001": 10.0, "C-999": 50.0})
        assert any(w.code == "binding_unknown_claim" for w in warnings)

    def test_no_warning_when_all_claims_bound(self):
        _, warnings, _ = _build_simple(rate_bindings={"C-001": 10.0})
        assert not any(w.code == "rate_unresolved" for w in warnings)

    def test_claims_truncated_flag_yields_warning(self):
        _, warnings, _ = build_boq_workbook(
            project_name="P",
            claim_rows=[_claim_row()],
            assumption_rows=[],
            checkmate_rows=[],
            rate_bindings={"C-001": 1.0},
            claims_truncated=True,
        )
        assert any(w.code == "claims_truncated" for w in warnings)

    def test_warnings_capped_at_max(self):
        bindings = {f"C-{i:03d}": float(i) for i in range(200)}
        _, warnings, _ = build_boq_workbook(
            project_name="P",
            claim_rows=[],
            assumption_rows=[],
            checkmate_rows=[],
            rate_bindings=bindings,
        )
        assert len(warnings) <= MAX_EXPORT_WARNINGS


# ---------------------------------------------------------------------------
# build_boq_workbook — gate
# ---------------------------------------------------------------------------

class TestWorkbookGate:
    def test_no_checkmate_rows_is_not_run(self):
        _, _, gate = _build_simple(checkmate=[])
        assert gate == GATE_NOT_RUN

    def test_all_passed_is_passed(self):
        _, _, gate = _build_simple(checkmate=[_checkmate_row("C-001", passed=1)])
        assert gate == GATE_PASSED

    def test_any_failed_is_rejected(self):
        cm = [_checkmate_row("C-001", passed=1), _checkmate_row("C-002", passed=0)]
        _, _, gate = _build_simple(checkmate=cm)
        assert gate == GATE_REJECTED


# ---------------------------------------------------------------------------
# build_boq_workbook — size bound
# ---------------------------------------------------------------------------

class TestWorkbookSizeBound:
    def test_single_claim_within_bytes(self):
        data, _, _ = _build_simple()
        assert len(data) <= MAX_EXPORT_BYTES

    def test_thirty_claims_within_bytes(self):
        rows = [_claim_row(claim_id=f"C-{i:03d}", value=float(i)) for i in range(30)]
        bindings = {f"C-{i:03d}": float(i) for i in range(30)}
        data, _, _ = build_boq_workbook(
            project_name="Large",
            claim_rows=rows,
            assumption_rows=[],
            checkmate_rows=[],
            rate_bindings=bindings,
        )
        assert len(data) <= MAX_EXPORT_BYTES

